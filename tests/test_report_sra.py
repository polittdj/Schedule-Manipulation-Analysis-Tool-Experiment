"""SRA-section coverage for the Excel and Word reports (reference-tool capability).

The Monte-Carlo Schedule Risk Analysis must appear in both reports when an
``SRAResult`` is supplied -- with finish percentiles traceable to the result
(H-DRIFT-1) and the parity-honesty caption present (the SRA method is reference
parity; the DEFAULT duration spread is a tool heuristic) -- and must be ABSENT
when no ``sra`` is passed, so existing single-schedule reports are unchanged.

A zero-variance three-point estimator (O == M == P == duration) makes the
Monte-Carlo finish deterministic, so the percentile day values are exact and the
assertions are non-flaky (H-VACUOUS-TEST mitigation).
"""

from __future__ import annotations

import datetime as dt

from docx.document import Document

from schedule_forensics.analysis import analyze_schedule
from schedule_forensics.report_excel import build_excel_workbook
from schedule_forensics.report_word import build_word_document
from schedule_forensics.schemas import Relation, Schedule, Task
from schedule_forensics.sra import SRAResult, run_sra

_START = dt.datetime(2025, 1, 6, 8)


def _chain() -> Schedule:
    """A 1->2 FS chain (480 + 1440 = finish 1920 min = 4.0 working days)."""
    return Schedule(
        name="Risk",
        project_start=_START,
        status_date=_START,
        tasks=(
            Task(unique_id=1, name="A", duration_minutes=480),
            Task(unique_id=2, name="B", duration_minutes=1440),
        ),
        relations=(Relation(predecessor_id=1, successor_id=2),),
    )


def _deterministic_sra(schedule: Schedule) -> SRAResult:
    """SRA with a zero-variance spread: every trial reproduces the CPM finish."""
    return run_sra(
        schedule,
        iterations=32,
        seed=1,
        three_point=lambda t: (
            float(t.duration_minutes),
            float(t.duration_minutes),
            float(t.duration_minutes),
        ),
    )


def _all_cell_values(wb: object) -> set[object]:
    sheet = wb["Risk (SRA)"]  # type: ignore[index]
    return {cell.value for row in sheet.iter_rows() for cell in row}


def _all_doc_text(doc: Document) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


# ── Excel ─────────────────────────────────────────────────────────────────────


def test_excel_sra_sheet_present_and_traceable() -> None:
    sched = _chain()
    wb = build_excel_workbook(analyze_schedule(sched), sra=_deterministic_sra(sched))
    assert "Risk (SRA)" in wb.sheetnames
    values = _all_cell_values(wb)
    # P50/P80/P95 labels rendered.
    assert {"P50", "P80", "P95"} <= {v for v in values if isinstance(v, str)}
    # Deterministic finish = 1920 min = 4.0 working days (openpyxl may store 4.0 as int 4).
    assert 4.0 in values or 4 in values
    # Both chain tasks are critical in every trial -> criticality index 100%.
    assert 100.0 in values or 100 in values
    assert 1 in values and 2 in values  # the two activity UniqueIDs


def test_excel_sra_parity_caption_present() -> None:
    sched = _chain()
    wb = build_excel_workbook(analyze_schedule(sched), sra=_deterministic_sra(sched))
    values = _all_cell_values(wb)
    blob = " ".join(str(v) for v in values)
    # Parity-honesty (LAW 2): the default spread is named a tool heuristic, not parity.
    assert "source-pending" in blob
    assert "Acumen" in blob


def test_excel_sra_sheet_absent_without_sra() -> None:
    sched = _chain()
    assert "Risk (SRA)" not in build_excel_workbook(analyze_schedule(sched)).sheetnames


# ── Word ────────────────────────────────────────────────────────────────────


def test_word_sra_section_present_and_traceable() -> None:
    sched = _chain()
    doc = build_word_document(analyze_schedule(sched), sra=_deterministic_sra(sched))
    text = _all_doc_text(doc)
    assert "Schedule Risk Analysis (Monte Carlo)" in text
    assert "P50" in text and "P80" in text and "P95" in text
    assert "4.0" in text  # deterministic finish in working days
    assert "100.0" in text  # criticality index for the (always-critical) chain
    assert "source-pending" in text  # parity-honesty caption


def test_word_sra_section_absent_without_sra() -> None:
    sched = _chain()
    doc = build_word_document(analyze_schedule(sched))
    assert "Schedule Risk Analysis" not in _all_doc_text(doc)
